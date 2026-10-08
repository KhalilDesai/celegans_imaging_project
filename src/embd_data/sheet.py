"""Convert the published ``genes`` table dump into a canonical, exactly round-trippable CSV.

Run with Python 3 (openpyxl). The canonical CSV is the only thing the legacy (Python 2.7)
baseline reads, so both interpreters see bit-identical float64 values:

* floats are written with ``repr`` (shortest round-trip repr in Python 3); Python 2.7's
  ``float()`` is correctly rounded, so it parses them back to the same double;
* ints are written as integers (they become floats in FLOAT columns, as in MySQL);
* empty cells and the string ``'NULL'`` become empty fields (SQL NULL / NaN);
* columns with an empty header (``Unnamed: 5`` in pandas) are dropped;
* columns the original code expects but the dump lacks are added (see DECISIONS A4, R3):
  ``id`` (= ``mysql_id``; the control row gets 2627, the id hardcoded in
  ``RNAiClass.write_controlAvg_mysql``), ``version``, ``prev_GLS``, ``prev_MS``.

Usage::

    python -m embd_data.sheet data/All_genes_mySQL_table.xlsx data/Index_EMBD_ID_Gene_names.xlsx OUTDIR
"""
import csv
import hashlib
import json
import math
import os
import sys

CONTROL_MYSQL_ID = 2627          # RNAiClass.write_controlAvg_mysql: "row number id for EMBD0000"
DATA_VERSION = 'zenodo-7761843'  # stands in for the git hash stored in genes.version (DECISIONS A4)


def _read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    body = [list(r) for r in rows]
    wb.close()
    return header, body


def _cell_to_text(v):
    if v is None:
        return ''
    if isinstance(v, bool):
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if math.isnan(v):
            return ''
        return repr(v)
    s = str(v)
    if s == 'NULL':
        return ''
    return s


def convert_genes_table(xlsx_path, out_csv):
    header, body = _read_xlsx(xlsx_path)
    keep = [i for i, h in enumerate(header) if h is not None and str(h).strip() != '']
    names = [str(header[i]) for i in keep]
    if len(set(names)) != len(names):
        raise ValueError('duplicate column names in sheet')
    i_mysql, i_embd = names.index('mysql_id'), names.index('embd_id')
    out_names = ['id'] + names + ['version', 'prev_GLS', 'prev_MS']
    n_null_strings = 0
    with open(out_csv, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(out_names)
        for r in body:
            vals = [r[i] for i in keep]
            if all(v is None for v in vals):
                continue
            n_null_strings += sum(1 for v in vals if v == 'NULL')
            mysql_id = vals[i_mysql]
            embd = int(vals[i_embd])
            if mysql_id is None:
                if embd != 0:
                    raise ValueError('missing mysql_id for embd %d' % embd)
                mysql_id = CONTROL_MYSQL_ID
            w.writerow([str(int(mysql_id))] + [_cell_to_text(v) for v in vals] + [DATA_VERSION, '', ''])
    return {'columns': len(out_names), 'rows': len(body), 'null_strings_converted': n_null_strings}


def write_gene_key(index_xlsx, out_csv):
    """Reconstruct ``gene_key.csv`` (read by embdFunc.getGeneName/getGeneCommonName).

    The original file is not published (DECISIONS A9). Layout expected by the code: header row, then
    ``EMB_P####, <sequence name>, <common name>``. It is rebuilt from the 2023 Zenodo index, so names that
    changed between 2022 and 2023 differ from the originals (DECISIONS A12).
    """
    header, body = _read_xlsx(index_xlsx)
    h = [str(x) for x in header]
    ie, ic, ig = h.index('EMBD ID'), h.index('C. elegans ID'), h.index('Gene Name')
    with open(out_csv, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(['EMBD_ID', 'Sequence_Name', 'Common_Name'])
        for r in body:
            if r[ie] is None:
                continue
            n = int(r[ie])
            if n == 0:
                continue
            w.writerow(['EMB_P{0:04}'.format(n), r[ic] or '', r[ig] or ''])


def _sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def main(argv):
    sheet, index, outdir = argv[1:4]
    os.makedirs(outdir, exist_ok=True)
    genes_csv = os.path.join(outdir, 'genes_table.csv')
    key_csv = os.path.join(outdir, 'gene_key.csv')
    info = convert_genes_table(sheet, genes_csv)
    write_gene_key(index, key_csv)
    info['inputs'] = {os.path.basename(p): _sha256(p) for p in (sheet, index)}
    info['outputs'] = {os.path.basename(p): _sha256(p) for p in (genes_csv, key_csv)}
    with open(os.path.join(outdir, 'conversion_manifest.json'), 'w') as f:
        json.dump(info, f, indent=2, sort_keys=True)
    print(json.dumps(info, indent=2, sort_keys=True))


if __name__ == '__main__':
    main(sys.argv)
